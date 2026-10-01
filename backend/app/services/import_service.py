import csv
import io
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook
from sqlalchemy import select

from app.core.config import get_settings
from app.models.business import Business
from app.models.category import Category
from app.models.transaction import Transaction
from app.models.uploaded_file import UploadedFile
from app.schemas.imports import (
    ColumnMapping,
    ConfirmRow,
    ParsedRow,
    ValidatedRow,
    ValidateResponse,
)

settings = get_settings()

MAX_ROWS = 5000

TARGET_FIELDS = ("transaction_date", "description", "amount", "type", "category", "payment_method")

FIELD_SYNONYMS = {
    "transaction_date": ["date", "transaction date", "transaction_date", "txn date", "posted date"],
    "description": ["description", "desc", "details", "memo", "narration", "particulars"],
    "amount": ["amount", "value", "total", "amount (rwf)"],
    "type": ["type", "transaction type", "transaction_type", "category type"],
    "category": ["category", "expense category", "income category"],
    "payment_method": ["payment method", "payment_method", "method", "channel"],
}

TYPE_SYNONYMS = {
    "income": "income",
    "revenue": "income",
    "credit": "income",
    "expense": "expense",
    "expenses": "expense",
    "cost": "expense",
    "debit": "expense",
    "purchase": "expense",
    "sale": "sale",
    "sales": "sale",
}

PAYMENT_METHOD_SYNONYMS = {
    "cash": "cash",
    "bank": "bank",
    "bank transfer": "bank",
    "transfer": "bank",
    "mobile money": "mobile_money",
    "mobile_money": "mobile_money",
    "momo": "mobile_money",
    "mtn": "mobile_money",
    "mtn mobile money": "mobile_money",
    "airtel money": "mobile_money",
    "card": "card",
    "credit card": "card",
    "debit card": "card",
    "visa": "card",
    "mastercard": "card",
}

DATE_FORMATS = ["%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y", "%d %b %Y", "%d %B %Y", "%B %d, %Y", "%b %d, %Y"]

CURRENCY_CHARS = re.compile(r"[^\d.\-()]")


class ImportError_(Exception):
    pass


def parse_file(filename: str, content: bytes) -> tuple[list[str], list[dict[str, str]]]:
    lower = filename.lower()
    if lower.endswith(".csv"):
        columns, rows = _parse_csv(content)
    elif lower.endswith(".xlsx") or lower.endswith(".xls"):
        columns, rows = _parse_xlsx(content)
    else:
        raise ImportError_("Unsupported file type. Please upload a .csv or .xlsx file.")

    if len(rows) > MAX_ROWS:
        raise ImportError_(f"This file has {len(rows)} rows; please import up to {MAX_ROWS} rows per file.")
    if not rows:
        raise ImportError_("This file has no data rows.")

    return columns, rows


def _parse_csv(content: bytes) -> tuple[list[str], list[dict[str, str]]]:
    text = content.decode("utf-8-sig", errors="replace")
    reader = csv.DictReader(io.StringIO(text))
    columns = [c.strip() for c in (reader.fieldnames or [])]
    rows = []
    for raw_row in reader:
        rows.append({(k or "").strip(): (v or "").strip() for k, v in raw_row.items() if k})
    return columns, rows


def _parse_xlsx(content: bytes) -> tuple[list[str], list[dict[str, str]]]:
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    rows_iter = sheet.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        return [], []

    columns = [str(c).strip() if c is not None else "" for c in header_row]
    rows = []
    for raw_row in rows_iter:
        if all(cell is None for cell in raw_row):
            continue
        row = {}
        for col, cell in zip(columns, raw_row):
            if not col:
                continue
            if isinstance(cell, (datetime, date)):
                row[col] = cell.isoformat()
            elif cell is None:
                row[col] = ""
            else:
                row[col] = str(cell).strip()
        rows.append(row)
    return columns, rows


def suggest_mapping(columns: list[str]) -> dict[str, str | None]:
    normalized = {c: _normalize_header(c) for c in columns}
    mapping: dict[str, str | None] = {}
    for field in TARGET_FIELDS:
        synonyms = FIELD_SYNONYMS[field]
        match = next((col for col, norm in normalized.items() if norm in synonyms), None)
        mapping[field] = match
    return mapping


def _normalize_header(header: str) -> str:
    return re.sub(r"[_\s]+", " ", header.strip().lower())


def _parse_amount(raw: str) -> Decimal | None:
    if not raw or not raw.strip():
        return None
    value = raw.strip()
    negative = value.startswith("(") and value.endswith(")")
    cleaned = CURRENCY_CHARS.sub("", value)
    cleaned = cleaned.replace("(", "").replace(")", "")
    if not cleaned or cleaned in ("-", "."):
        return None
    try:
        amount = Decimal(cleaned)
    except InvalidOperation:
        return None
    amount = abs(amount)
    return amount if amount > 0 else None


def _parse_date(raw: str) -> date | None:
    if not raw or not raw.strip():
        return None
    value = raw.strip()
    try:
        return date.fromisoformat(value[:10])
    except ValueError:
        pass
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(value, fmt).date()
        except ValueError:
            continue
    return None


def _normalize_type(raw: str | None, amount_sign_negative: bool) -> tuple[str, bool]:
    if raw and raw.strip():
        key = raw.strip().lower()
        if key in TYPE_SYNONYMS:
            return TYPE_SYNONYMS[key], False
        return "expense" if amount_sign_negative else "income", True
    return "expense" if amount_sign_negative else "income", False


def _normalize_payment_method(raw: str | None) -> tuple[str, bool]:
    if raw and raw.strip():
        key = raw.strip().lower()
        if key in PAYMENT_METHOD_SYNONYMS:
            return PAYMENT_METHOD_SYNONYMS[key], False
        return "other", True
    return "other", False


def _match_category(categories: list[Category], category_type: str, raw_name: str | None) -> Category | None:
    if not raw_name or not raw_name.strip():
        return None
    target = raw_name.strip().lower()
    candidates = [c for c in categories if c.type == category_type]
    for c in candidates:
        if c.name.lower() == target:
            return c
    for c in candidates:
        if target in c.name.lower() or c.name.lower() in target:
            return c
    return None


def _has_sign(raw: str) -> bool:
    return raw.strip().startswith("-") or (raw.strip().startswith("(") and raw.strip().endswith(")"))


def validate_rows(db, business: Business, rows: list[ParsedRow], mapping: ColumnMapping) -> ValidateResponse:
    categories = list(db.scalars(select(Category).where(Category.business_id == business.id)))

    existing_keys = {
        (t.transaction_date, t.amount, t.description.strip().lower())
        for t in db.scalars(select(Transaction).where(Transaction.business_id == business.id))
    }
    seen_in_batch: set[tuple] = set()

    results: list[ValidatedRow] = []
    needs_ai: list[dict] = []

    for row in rows:
        issues: list[str] = []
        raw = row.raw

        amount_raw = raw.get(mapping.amount, "") if mapping.amount else ""
        amount = _parse_amount(amount_raw)
        if amount is None:
            issues.append("Missing or invalid amount")

        date_raw = raw.get(mapping.transaction_date, "") if mapping.transaction_date else ""
        transaction_date = _parse_date(date_raw)
        if transaction_date is None:
            issues.append("Missing or invalid date")

        description = (raw.get(mapping.description, "") if mapping.description else "").strip()
        if not description:
            description = "Imported transaction"
            issues.append("Missing description -- used a placeholder")

        type_raw = raw.get(mapping.type, "") if mapping.type else None
        txn_type, type_was_guessed = _normalize_type(type_raw, _has_sign(amount_raw))
        if type_was_guessed and type_raw:
            issues.append(f"Unrecognized transaction type '{type_raw}' -- inferred from amount")

        payment_raw = raw.get(mapping.payment_method, "") if mapping.payment_method else None
        payment_method, payment_was_guessed = _normalize_payment_method(payment_raw)
        if payment_was_guessed:
            issues.append(f"Unrecognized payment method '{payment_raw}' -- defaulted to Other")

        category_bucket = "expense" if txn_type == "expense" else "income"
        category_raw = raw.get(mapping.category, "") if mapping.category else None
        category = _match_category(categories, category_bucket, category_raw)
        if category_raw and not category:
            issues.append(f"Category '{category_raw}' not found -- will need review")
        elif not category:
            needs_ai.append(
                {"row_number": row.row_number, "description": description, "type": txn_type, "bucket": category_bucket}
            )

        if amount is not None and transaction_date is not None:
            dedupe_key = (transaction_date, amount, description.strip().lower())
            if dedupe_key in existing_keys or dedupe_key in seen_in_batch:
                issues.append("Possible duplicate transaction")
            seen_in_batch.add(dedupe_key)

        if amount is None or transaction_date is None:
            status = "invalid"
        elif issues:
            status = "warning"
        else:
            status = "valid"

        results.append(
            ValidatedRow(
                row_number=row.row_number,
                status=status,
                issues=issues,
                transaction_date=transaction_date,
                description=description,
                amount=amount,
                type=txn_type,
                payment_method=payment_method,
                category_id=str(category.id) if category else None,
                category_name=category.name if category else None,
            )
        )

    ai_available = _suggest_ai_categories(business, categories, needs_ai, results)

    valid_count = sum(1 for r in results if r.status == "valid")
    warning_count = sum(1 for r in results if r.status == "warning")
    invalid_count = sum(1 for r in results if r.status == "invalid")

    return ValidateResponse(
        valid_count=valid_count,
        warning_count=warning_count,
        invalid_count=invalid_count,
        rows=results,
        ai_categorization_available=ai_available,
    )


def _suggest_ai_categories(
    business: Business, categories: list[Category], needs_ai: list[dict], results: list[ValidatedRow]
) -> bool:
    if not settings.gemini_api_key or not needs_ai:
        return bool(settings.gemini_api_key)

    try:
        from google import genai
        from google.genai import types

        income_names = [c.name for c in categories if c.type == "income"]
        expense_names = [c.name for c in categories if c.type == "expense"]

        items_desc = "\n".join(
            f"{item['row_number']}. ({item['bucket']}) {item['description']}" for item in needs_ai[:200]
        )
        prompt = (
            f"Available income categories: {', '.join(income_names)}\n"
            f"Available expense categories: {', '.join(expense_names)}\n\n"
            "For each numbered transaction below, pick the single best-matching category from the "
            "appropriate list (income transactions get an income category, expense transactions get an "
            "expense category). Respond as a JSON array of objects with fields: row_number (int), "
            "category_name (string, must exactly match one of the provided category names), and "
            "confidence (integer 0-100).\n\n"
            f"Transactions:\n{items_desc}"
        )

        client = genai.Client(api_key=settings.gemini_api_key)
        response = client.models.generate_content(
            model="gemini-3.1-flash-lite",
            contents=prompt,
            config=types.GenerateContentConfig(response_mime_type="application/json"),
        )

        import json

        suggestions = json.loads(response.text)
        by_row = {int(s["row_number"]): s for s in suggestions if "row_number" in s}
        by_name = {c.name.lower(): c for c in categories}

        for row in results:
            suggestion = by_row.get(row.row_number)
            if not suggestion:
                continue
            match = by_name.get(str(suggestion.get("category_name", "")).lower())
            if match:
                row.suggested_category_id = str(match.id)
                row.suggested_category_name = match.name
                row.suggested_category_confidence = float(suggestion.get("confidence", 0))
        return True
    except Exception:  # noqa: BLE001 -- AI categorization is a best-effort enhancement, never a hard requirement
        return False


def confirm_import(db, business: Business, filename: str, rows: list[ConfirmRow]) -> int:
    for row in rows:
        db.add(
            Transaction(
                business_id=business.id,
                category_id=row.category_id,
                type=row.type,
                amount=row.amount,
                description=row.description,
                transaction_date=row.transaction_date,
                payment_method=row.payment_method,
                source="import",
            )
        )

    db.add(
        UploadedFile(
            business_id=business.id,
            filename=filename,
            row_count=len(rows),
            imported_count=len(rows),
        )
    )
    db.commit()
    return len(rows)
