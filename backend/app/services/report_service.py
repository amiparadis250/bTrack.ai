import io
import uuid
from datetime import date, datetime, timezone
from decimal import Decimal

from fastapi import HTTPException, status
from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy.orm import Session

from app.models.business import Business
from app.models.category import Category
from app.models.report import Report
from app.models.transaction import Transaction
from app.repositories import transaction_repository
from app.services import analytics_service

CONTENT_TYPES = {
    "pdf": "application/pdf",
    "excel": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}
FILE_EXTENSIONS = {"pdf": "pdf", "excel": "xlsx"}

REPORT_TITLES = {
    "financial_summary": "Financial Summary",
    "transactions": "Transaction Report",
    "sales": "Sales Report",
    "expenses": "Expense Report",
}

_TYPE_FILTER = {"financial_summary": None, "transactions": None, "sales": "sale", "expenses": "expense"}


def create_report(db: Session, business: Business, report_type: str, format: str, start: date, end: date) -> Report:
    report = Report(
        business_id=business.id, report_type=report_type, format=format, period_start=start, period_end=end
    )
    db.add(report)
    db.commit()
    db.refresh(report)
    return report


def get_report_or_404(db: Session, business: Business, report_id: uuid.UUID) -> Report:
    report = db.get(Report, report_id)
    if not report or report.business_id != business.id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Report not found.")
    return report


def _category_names(db: Session, business: Business) -> dict[uuid.UUID, str]:
    categories = db.query(Category).filter(Category.business_id == business.id).all()
    return {c.id: c.name for c in categories}


def _transactions_for(db: Session, business: Business, report: Report) -> list[Transaction]:
    rows, _total = transaction_repository.list_transactions(
        db,
        business.id,
        type_=_TYPE_FILTER[report.report_type],
        date_from=report.period_start,
        date_to=report.period_end,
        limit=10_000,
        offset=0,
    )
    return rows


def build_report_file(db: Session, business: Business, report: Report) -> tuple[bytes, str]:
    safe_name = "".join(c if c.isalnum() else "_" for c in business.name)
    filename = (
        f"{safe_name}_{report.report_type}_{report.period_start}_{report.period_end}."
        f"{FILE_EXTENSIONS[report.format]}"
    )
    content = _build_pdf(db, business, report) if report.format == "pdf" else _build_excel(db, business, report)
    return content, filename


def _styled_table(data: list[list[str]]) -> Table:
    table = Table(data, hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0d9488")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#d6dbdd")),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f7f9f9")]),
            ]
        )
    )
    return table


def _build_pdf(db: Session, business: Business, report: Report) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=2 * cm, bottomMargin=2 * cm)
    styles = getSampleStyleSheet()
    elements = [
        Paragraph(business.name, styles["Title"]),
        Paragraph(REPORT_TITLES[report.report_type], styles["Heading2"]),
        Paragraph(f"Period: {report.period_start} to {report.period_end}", styles["Normal"]),
        Paragraph(f"Generated: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}", styles["Normal"]),
        Spacer(1, 0.5 * cm),
    ]

    if report.report_type == "financial_summary":
        overview = analytics_service.get_overview(db, business.id, report.period_start, report.period_end)
        margin = f"{overview.profit_margin:.1f}%" if overview.profit_margin is not None else "n/a"
        data = [
            ["Metric", "Amount"],
            ["Revenue", f"{business.currency} {overview.revenue}"],
            ["Expenses", f"{business.currency} {overview.expenses}"],
            ["Profit", f"{business.currency} {overview.profit}"],
            ["Profit Margin", margin],
            ["Cash Inflow", f"{business.currency} {overview.cash_inflow}"],
            ["Cash Outflow", f"{business.currency} {overview.cash_outflow}"],
            ["Net Cash Flow", f"{business.currency} {overview.net_cash_flow}"],
        ]
        elements.append(_styled_table(data))

        breakdown = analytics_service.get_expense_breakdown(db, business.id, report.period_start, report.period_end)
        if breakdown:
            elements.append(Spacer(1, 0.6 * cm))
            elements.append(Paragraph("Expense Breakdown", styles["Heading3"]))
            bdata = [["Category", "Amount", "% of Expenses"]] + [
                [item.category_name, f"{business.currency} {item.total}", f"{item.percentage:.0f}%"]
                for item in breakdown
            ]
            elements.append(_styled_table(bdata))
    else:
        category_names = _category_names(db, business)
        rows = _transactions_for(db, business, report)
        data = [["Date", "Description", "Category", "Type", "Payment Method", "Amount"]]
        net = Decimal("0")
        for t in rows:
            data.append(
                [
                    str(t.transaction_date),
                    t.description,
                    category_names.get(t.category_id, "Uncategorized") if t.category_id else "Uncategorized",
                    t.type,
                    t.payment_method,
                    f"{business.currency} {t.amount}",
                ]
            )
            net += t.amount if t.type != "expense" else -t.amount
        if len(data) == 1:
            elements.append(Paragraph("No transactions recorded for this period.", styles["Normal"]))
        else:
            elements.append(_styled_table(data))
            elements.append(Spacer(1, 0.4 * cm))
            elements.append(Paragraph(f"Net total: {business.currency} {net}", styles["Heading4"]))

    doc.build(elements)
    return buffer.getvalue()


def _build_excel(db: Session, business: Business, report: Report) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append([business.name])
    summary.append([REPORT_TITLES[report.report_type]])
    summary.append([f"Period: {report.period_start} to {report.period_end}"])
    summary.append([f"Generated: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}"])
    summary.append([])

    if report.report_type == "financial_summary":
        overview = analytics_service.get_overview(db, business.id, report.period_start, report.period_end)
        summary.append(["Metric", "Amount"])
        summary.append(["Revenue", str(overview.revenue)])
        summary.append(["Expenses", str(overview.expenses)])
        summary.append(["Profit", str(overview.profit)])
        summary.append(["Profit Margin %", overview.profit_margin])
        summary.append(["Cash Inflow", str(overview.cash_inflow)])
        summary.append(["Cash Outflow", str(overview.cash_outflow)])
        summary.append(["Net Cash Flow", str(overview.net_cash_flow)])

        breakdown = analytics_service.get_expense_breakdown(db, business.id, report.period_start, report.period_end)
        if breakdown:
            sheet = workbook.create_sheet("Expense Breakdown")
            sheet.append(["Category", "Amount", "Percentage"])
            for item in breakdown:
                sheet.append([item.category_name, str(item.total), item.percentage])
    else:
        category_names = _category_names(db, business)
        rows = _transactions_for(db, business, report)
        sheet = workbook.create_sheet("Transactions")
        sheet.append(["Date", "Description", "Category", "Type", "Payment Method", "Amount"])
        for t in rows:
            sheet.append(
                [
                    str(t.transaction_date),
                    t.description,
                    category_names.get(t.category_id, "Uncategorized") if t.category_id else "Uncategorized",
                    t.type,
                    t.payment_method,
                    str(t.amount),
                ]
            )

    for sheet in workbook.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
